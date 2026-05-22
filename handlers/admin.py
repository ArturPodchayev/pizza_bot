import asyncio
import logging
from io import BytesIO
from datetime import datetime

from aiogram import Router, Bot, F
from aiogram.filters import Command
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.types import Message, InlineKeyboardMarkup, InlineKeyboardButton, CallbackQuery, BufferedInputFile
from aiogram.exceptions import TelegramForbiddenError, TelegramBadRequest

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

import db
from config import ADMIN_IDS

router = Router()
logger = logging.getLogger(__name__)

BROADCAST_TARGETS = {
    "broadcast_all": ("all", "Все / All / Hammasi"),
    "broadcast_ru":  ("ru",  "🇷🇺 Только RU"),
    "broadcast_uz":  ("uz",  "🇺🇿 Только UZ"),
    "broadcast_en":  ("en",  "🇬🇧 Только EN"),
}


class Broadcast(StatesGroup):
    choosing_target = State()
    waiting_text = State()


def is_admin(user_id: int) -> bool:
    return user_id in ADMIN_IDS


def broadcast_target_keyboard() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [InlineKeyboardButton(text="📢 Все участники", callback_data="broadcast_all")],
            [
                InlineKeyboardButton(text="🇷🇺 RU", callback_data="broadcast_ru"),
                InlineKeyboardButton(text="🇺🇿 UZ", callback_data="broadcast_uz"),
                InlineKeyboardButton(text="🇬🇧 EN", callback_data="broadcast_en"),
            ],
        ]
    )


def make_excel(users: list) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Participants"

    headers = ["#", "Telegram ID", "Username", "Language", "Full Name", "Phone", "Registered At", "Ref Code"]
    header_fill = PatternFill("solid", fgColor="F7931A")  # Bitcoin orange
    header_font = Font(bold=True, color="FFFFFF")

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 30
    ws.column_dimensions["F"].width = 20
    ws.column_dimensions["G"].width = 20
    ws.column_dimensions["H"].width = 20

    for i, u in enumerate(users, 1):
        ws.append([
            i,
            u["telegram_id"],
            u["username"] or "—",
            u["language"].upper() if u["language"] else "—",
            u["full_name"],
            u["phone"],
            u["registered_at"],
            u["ref_code"] or "—",
        ])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@router.message(Command("admin"))
async def cmd_admin(message: Message) -> None:
    if not is_admin(message.from_user.id):
        return

    count = await db.get_users_count()
    ru = len(await db.get_users_by_language("ru"))
    uz = len(await db.get_users_by_language("uz"))
    en = len(await db.get_users_by_language("en"))

    await message.answer(
        f"🍕 <b>Bitcoin Pizza Fest 2026 — Админ-панель</b>\n\n"
        f"👥 Всего участников: <b>{count}</b>\n"
        f"  🇷🇺 RU: {ru}\n"
        f"  🇺🇿 UZ: {uz}\n"
        f"  🇬🇧 EN: {en}\n\n"
        f"<b>Команды:</b>\n"
        f"/export — скачать базу в Excel\n"
        f"/broadcast — сделать рассылку\n"
        f"/users — количество участников\n"
        f"/refs — реферальная статистика\n"
        f"/export_humo — база квалифицированных HUMO",
        parse_mode="HTML",
    )


@router.message(Command("users"))
async def cmd_users(message: Message) -> None:
    if not is_admin(message.from_user.id):
        return
    count = await db.get_users_count()
    await message.answer(f"👥 Всего участников: <b>{count}</b>", parse_mode="HTML")


def make_humo_excel(users: list) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "HUMO Qualified"

    headers = ["#", "Telegram ID", "Username", "Full Name", "Phone", "Score", "Completed At", "Status"]
    header_fill = PatternFill("solid", fgColor="F7931A")
    header_font = Font(bold=True, color="FFFFFF")

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 30
    ws.column_dimensions["E"].width = 20
    ws.column_dimensions["F"].width = 8
    ws.column_dimensions["G"].width = 22
    ws.column_dimensions["H"].width = 12

    for i, u in enumerate(users, 1):
        ws.append([
            i,
            u["telegram_id"],
            u["username"] or "—",
            u["full_name"],
            u["phone"],
            u["score"],
            str(u["completed_at"]),
            "qualified",
        ])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@router.message(Command("refs"))
async def cmd_refs(message: Message) -> None:
    if not is_admin(message.from_user.id):
        return

    rows = await db.get_ref_stats()
    direct_count = await db.get_direct_users_count()

    lines = ["📊 <b>Реферальная статистика:</b>\n"]
    for row in rows:
        lines.append(f"{row['ref_code']} — {row['cnt']} чел.")
    lines.append(f"(прямые) — {direct_count} чел.")

    await message.answer("\n".join(lines), parse_mode="HTML")


@router.message(Command("export"))
async def cmd_export(message: Message) -> None:
    if not is_admin(message.from_user.id):
        return

    users = await db.get_all_active_users()
    if not users:
        await message.answer("База пока пуста.")
        return

    buf = make_excel(users)
    filename = f"bpf2026_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    file = BufferedInputFile(buf.read(), filename=filename)

    await message.answer_document(
        file,
        caption=f"📥 База участников\n👥 Всего: {len(users)}",
    )


@router.message(Command("export_humo"))
async def cmd_export_humo(message: Message) -> None:
    if not is_admin(message.from_user.id):
        return

    users = await db.get_qualified_users()
    if not users:
        await message.answer("Нет квалифицированных участников HUMO.")
        return

    buf = make_humo_excel(users)
    filename = f"humo_qualified_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    file = BufferedInputFile(buf.read(), filename=filename)

    await message.answer_document(
        file,
        caption=f"📥 HUMO: квалифицированные участники\n✅ Всего: {len(users)}",
    )


@router.message(Command("broadcast"))
async def cmd_broadcast(message: Message, state: FSMContext) -> None:
    if not is_admin(message.from_user.id):
        return

    await message.answer(
        "Выберите аудиторию для рассылки:",
        reply_markup=broadcast_target_keyboard(),
    )
    await state.set_state(Broadcast.choosing_target)


@router.callback_query(Broadcast.choosing_target, F.data.startswith("broadcast_"))
async def cb_broadcast_target(callback: CallbackQuery, state: FSMContext) -> None:
    target_key = callback.data
    target_info = BROADCAST_TARGETS.get(target_key)
    if not target_info:
        return

    target, label = target_info
    await state.update_data(broadcast_target=target, broadcast_label=label)
    await callback.message.edit_text(
        f"Рассылка для: <b>{label}</b>\n\nВведите текст сообщения:",
        parse_mode="HTML",
    )
    await callback.answer()
    await state.set_state(Broadcast.waiting_text)


@router.message(Broadcast.waiting_text)
async def do_broadcast(message: Message, state: FSMContext, bot: Bot) -> None:
    if not is_admin(message.from_user.id):
        return

    data = await state.get_data()
    target = data.get("broadcast_target", "all")
    label = data.get("broadcast_label", "Все")

    if target == "all":
        users = await db.get_all_active_users()
    else:
        users = await db.get_users_by_language(target)

    await state.clear()
    status_msg = await message.answer(f"⏳ Начинаю рассылку для {label}...")

    sent, failed = 0, 0
    for user in users:
        try:
            await bot.send_message(user["telegram_id"], message.text)
            sent += 1
        except (TelegramForbiddenError, TelegramBadRequest) as e:
            logger.warning(f"Broadcast failed for {user['telegram_id']}: {e}")
            await db.deactivate_user(user["telegram_id"])
            failed += 1
        except Exception as e:
            logger.warning(f"Broadcast error for {user['telegram_id']}: {e}")
            failed += 1
        await asyncio.sleep(0.05)  # не флудить Telegram API

    await status_msg.edit_text(
        f"✅ Рассылка завершена — {label}\n\n"
        f"✅ Отправлено: {sent}\n"
        f"❌ Не доставлено: {failed}",
    )
